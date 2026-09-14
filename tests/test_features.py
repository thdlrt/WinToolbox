"""All tests use temporary data and injected providers; no user configuration access."""
import contextlib
import hashlib
import json
import os
import sqlite3
import sys
import threading
import uuid
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "core"))

from toolbox.features import backups, codex, files, knowledge, plugins
from toolbox.features._common import extract_zip, write_json


class Settings:
    def __init__(self, directory):
        self.directory = directory
        self.secrets = {"example": "secret-value-test-only"}
        self.reload()

    def reload(self):
        path = self.directory / "settings.json"
        self.value = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"preferences": {}, "roles": {}}

    def get(self):
        return self.value

    def export_secrets(self):
        return dict(self.secrets)

    def import_secrets(self, values):
        self.secrets = dict(values)
        write_json(self.directory / "secrets.json", {"protected-test": values})


class Providers:
    def __init__(self):
        self.messages = []
        self.fail_embeddings = False

    def chat(self, messages, role="chat", stream_callback=None):
        self.messages.append((role, messages))
        if role == "translate":
            request = json.loads(messages[-1]["content"])
            return json.dumps({f["id"]: "翻译 " + f["name"] for f in request["files"]}, ensure_ascii=False)
        if role == "vision":
            return "图表：准确率 92%，时间 2026 年。"
        return "资料说明转写延迟为三秒。[S1]"

    def embed(self, texts):
        if self.fail_embeddings:
            raise ValueError("测试嵌入未配置")
        return [[1.0, float("转写" in t), float("latency" in t.lower()), 0.2] for t in texts]


class Job:
    def __init__(self, params):
        self.params = params
        self.id = uuid.uuid4().hex
        self.artifacts = []
        self.cancelled = False

    def check_cancelled(self):
        if self.cancelled:
            raise RuntimeError("cancelled")

    def progress(self, percent, message):
        self.check_cancelled()

    def artifact(self, path, kind, label):
        assert Path(path).is_file()
        self.artifacts.append({"path": str(path), "kind": kind})


class Jobs:
    def __init__(self):
        self.runners = {}
        self.lock = threading.RLock()
        self.active = {}
        self.last_params = None

    def register(self, name, runner):
        self.runners[name] = runner

    def submit(self, tool, params, runner):
        self.last_params = params
        job = Job(params)
        result = runner(job)
        return {"id": job.id, "status": "completed", "result": result, "artifacts": job.artifacts}


class App:
    def __init__(self, directory):
        self.data_dir = directory
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.settings = Settings(directory)
        self.providers = Providers()
        self.jobs = Jobs()
        self.handlers = {}
        self.events = []

    def register(self, name, handler):
        self.handlers[name] = handler

    def emit(self, event, **kwargs):
        self.events.append((event, kwargs))

    def call(self, method, **params):
        return self.handlers[method](params)


@pytest.fixture
def app(tmp_path):
    return App(tmp_path / "app")


@pytest.fixture
def codex_home(tmp_path):
    directory = tmp_path / "isolated-codex"
    directory.mkdir()
    (directory / "models_cache.json").write_text(json.dumps({"models": [
        {"slug": "model-a", "display_name": "Model A", "context_window": 272000, "max_context_window": 1000000},
        {"slug": "unknown-max", "context_window": 1000},
        {"slug": "hidden", "visibility": "hide"}
    ]}), encoding="utf-8")
    (directory / "config.toml").write_bytes(b'# keep this\r\nmodel = "old" # selected\r\nmodel_auto_compact_token_limit = 500\r\n[profiles.test]\r\nmodel = "nested"\r\n')
    return directory


def test_codex_preserves_comments_tables_and_compact(codex_home):
    params = {"home": str(codex_home), "model": "model-a", "context": "500k"}
    before = (codex_home / "config.toml").read_bytes()
    preview = codex.preview(params)
    assert '# keep this\r\n' in preview["after"]
    assert 'model = "nested"' in preview["after"]
    assert 'model_auto_compact_token_limit = 500' in preview["after"]
    result = codex.apply({**params, "expected_hash": preview["hash"]})
    assert Path(result["backup_path"]).read_bytes() == before
    assert codex.scan(params)["current"]["context"] == 500000
    restored = codex.restore({**params, "backup_path": result["backup_path"], "expected_hash": result["hash"]})
    assert restored["ok"] and (codex_home / "config.toml").read_bytes() == before


def test_codex_rejects_external_change_and_overflow(codex_home):
    params = {"home": str(codex_home), "model": "model-a", "context": "max"}
    preview = codex.preview(params)
    with (codex_home / "config.toml").open("a", encoding="utf-8") as file:
        file.write("# concurrent change\n")
    with pytest.raises(ValueError, match="外部修改"):
        codex.apply({**params, "expected_hash": preview["hash"]})
    with pytest.raises(ValueError, match="超过"):
        codex.preview({**params, "context": "2m"})
    with pytest.raises(ValueError, match="未提供"):
        codex.preview({**params, "model": "unknown-max"})
    assert len(codex.scan(params)["models"]) == 2


def test_codex_corrupt_cache_and_auth_guard(codex_home):
    (codex_home / "models_cache.json").write_text("broken", encoding="utf-8")
    with pytest.raises(ValueError, match="无法解析"):
        codex.scan({"home": str(codex_home)})
    with pytest.raises(ValueError, match="认证"):
        codex.scan({"home": str(codex_home), "cache_path": str(codex_home / "auth.json")})


def test_codex_missing_file_then_bom(codex_home):
    config = codex_home / "config.toml"
    config.unlink()
    params = {"home": str(codex_home), "model": "model-a", "context": "default"}
    change = codex.preview(params)
    result = codex.apply({**params, "expected_hash": change["hash"]})
    assert result["backup_path"] is None
    config.write_bytes(b'\xef\xbb\xbfmodel="old"\n')
    change = codex.preview(params)
    codex.apply({**params, "expected_hash": change["hash"]})
    assert config.read_bytes().startswith(b"\xef\xbb\xbf")


def test_files_move_preview_revalidate_and_no_overwrite(app, tmp_path):
    files.register(app)
    source, dest = tmp_path / "输入", tmp_path / "输出"
    source.mkdir()
    (source / "one.mp4").write_bytes(b"original")
    (source / "other.txt").write_text("keep", encoding="utf-8")
    plan = app.call("files.preview", action="move", source_dir=str(source), dest_dir=str(dest), suffix=".mp4")
    assert len(plan["operations"]) == 1 and (source / "one.mp4").exists()
    dest.mkdir()
    (dest / "one.mp4").write_bytes(b"existing")
    with pytest.raises(ValueError, match="目标已存在"):
        app.call("files.apply", plan_id=plan["plan_id"])
    assert (dest / "one.mp4").read_bytes() == b"existing"
    (dest / "one.mp4").unlink()
    result = app.call("files.apply", plan_id=plan["plan_id"])
    assert result["moved"] == 1 and (source / "other.txt").exists()
    assert (dest / "one.mp4").read_bytes() == b"original"
    with pytest.raises(ValueError, match="已执行"):
        app.call("files.apply", plan_id=plan["plan_id"])


def test_files_source_modified_and_translation(app, tmp_path):
    files.register(app)
    source = tmp_path / "files"
    source.mkdir()
    file = source / "01-example.txt"
    file.write_text("a")
    plan = app.call("files.preview", action="translate", source_dir=str(source))
    assert plan["operations"][0]["target"].endswith("翻译 01-example.txt")
    file.write_text("changed")
    with pytest.raises(ValueError, match="已变化"):
        app.call("files.apply", plan_id=plan["plan_id"])
    assert file.read_text() == "changed"


def test_files_rollback_completed_entries(app, tmp_path, monkeypatch):
    files.register(app)
    source, dest = tmp_path / "from", tmp_path / "to"
    source.mkdir()
    for name in ("a.txt", "b.txt"):
        (source / name).write_text(name)
    plan = app.call("files.preview", action="move", source_dir=str(source), dest_dir=str(dest))
    original = files.move_exclusive
    def fail_second(src, dst):
        if Path(src).name == "b.txt":
            raise OSError("simulated I/O failure")
        return original(src, dst)
    monkeypatch.setattr(files, "move_exclusive", fail_second)
    with pytest.raises(ValueError, match="回滚"):
        app.call("files.apply", plan_id=plan["plan_id"])
    assert (source / "a.txt").read_text() == "a.txt" and (source / "b.txt").exists()
    assert not (dest / "a.txt").exists()


def test_files_prefix_translation_preserves_suffix_and_ignores_move_destination(app, tmp_path):
    files.register(app)
    source = tmp_path / "names"
    source.mkdir()
    (source / "Lecture_part_02.mp4").write_bytes(b"video")
    plan = app.call("files.preview", action="translate", source_dir=str(source), dest_dir=str(tmp_path / "stale-move-destination"), prefix_only=True)
    assert plan["operations"][0]["target"] == str(source / "翻译 Lecture_part_02.mp4")
    messages = app.providers.messages[-1][1]
    assert json.loads(messages[-1]["content"])["files"][0]["name"] == "Lecture"


@pytest.mark.parametrize("name", ["../escape.txt", "/rooted.txt", "C:/escape.txt", "a\\b.txt", "folder/../../escape", "CON.txt", "trailing. "])
def test_archive_rejects_windows_and_traversal_paths(tmp_path, name):
    archive = tmp_path / "bad.zip"
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.writestr(name.replace("\\", "/"), "evil")
    if "\\" in name:
        # zipfile normalizes native separators on Windows; patch both ZIP names.
        archive.write_bytes(archive.read_bytes().replace(name.replace("\\", "/").encode(), name.encode()))
    with zipfile.ZipFile(archive) as bundle, pytest.raises(ValueError):
        extract_zip(bundle, tmp_path / "extracted")
    assert not (tmp_path / "escape.txt").exists()


def test_archive_rejects_symlink_and_case_collision(tmp_path):
    for symlink in (False, True):
        archive = tmp_path / ("bad-link.zip" if symlink else "bad-case.zip")
        with zipfile.ZipFile(archive, "w") as bundle:
            if symlink:
                info = zipfile.ZipInfo("link")
                info.create_system = 3
                info.external_attr = 0o120777 << 16
                bundle.writestr(info, "../escape")
            else:
                bundle.writestr("FILE.py", "one")
                bundle.writestr("file.py", "two")
        with zipfile.ZipFile(archive) as bundle, pytest.raises(ValueError):
            extract_zip(bundle, tmp_path / "extracted")


def test_plugin_real_package_install_run_disable_uninstall(app, tmp_path):
    plugins.register(app)
    package = Path(__file__).parents[1] / "examples" / "file-hash.toolpkg"
    installed = app.call("plugins.install", path=str(package))
    assert installed["plugin"]["id"] == "file-hash"
    source = tmp_path / "source.dat"
    source.write_bytes(b"known source")
    result = app.call("plugins.run", id="file-hash", params={"path": str(source)})
    assert result["result"]["hash"] == hashlib.sha256(b"known source").hexdigest()
    app.call("plugins.enable", id="file-hash", enabled=False)
    with pytest.raises(ValueError, match="启用"):
        app.call("plugins.run", id="file-hash", params={"path": str(source)})
    app.call("plugins.uninstall", id="file-hash")
    assert app.call("plugins.list")["plugins"] == []


def test_plugin_bad_manifest_version_leaves_existing(app, tmp_path):
    plugins.register(app)
    package = Path(__file__).parents[1] / "examples" / "file-hash.toolpkg"
    app.call("plugins.install", path=str(package))
    bad = tmp_path / "new.toolpkg"
    with zipfile.ZipFile(package) as original, zipfile.ZipFile(bad, "w") as bundle:
        manifest = json.loads(original.read("manifest.json"))
        manifest["api_version"] = 9
        bundle.writestr("manifest.json", json.dumps(manifest))
        bundle.writestr("main.py", original.read("main.py"))
    with pytest.raises(ValueError, match="API 1"):
        app.call("plugins.install", path=str(bad))
    assert app.call("plugins.list")["plugins"][0]["version"] == "1.0.0"


def test_knowledge_text_sources_search_ask_and_cache(app, tmp_path):
    knowledge.register(app)
    collection = app.call("knowledge.create", name="汇报资料", prompt="回答要简短")
    source = tmp_path / "资料.md"
    source.write_text("# 系统\n转写延迟为三秒。\nThe latency is three seconds.\n", encoding="utf-8")
    result = app.call("knowledge.ingest", collection_id=collection["id"], paths=[str(source)], visual=False)
    assert result["result"]["documents"][0]["status"] == "ready"
    query = app.call("knowledge.search", collection_id=collection["id"], query="转写延迟")
    assert query["mode"] == "hybrid"
    assert "转写延迟为三秒" in query["results"][0]["quote"]
    assert query["results"][0]["page"] is None
    answer = app.call("knowledge.ask", collection_id=collection["id"], question="转写延迟是多少")
    assert "[S1]" in answer["answer"] and answer["sources"][0]["citation"] == "S1"
    cached = app.call("knowledge.ingest", collection_id=collection["id"], paths=[str(source)], visual=False)
    assert cached["result"]["documents"][0]["cached"]
    assert app.call("knowledge.documents", collection_id=collection["id"])["documents"][0]["pages"] == 1
    app.call("knowledge.delete", collection_id=collection["id"])
    assert not app.call("knowledge.list")["collections"]


def test_knowledge_keyword_fallback_is_explicit(app, tmp_path):
    knowledge.register(app)
    app.providers.fail_embeddings = True
    collection = app.call("knowledge.create", name="测试")
    source = tmp_path / "notes.txt"
    source.write_text("转写能力需要麦克风。", encoding="utf-8")
    ingestion = app.call("knowledge.ingest", collection_id=collection["id"], paths=[str(source)])
    assert ingestion["result"]["documents"][0]["status"] == "partial"
    query = app.call("knowledge.search", collection_id=collection["id"], query="转写")
    assert query["mode"] == "keyword" and query["warnings"] and query["results"]


def test_knowledge_pptx_slide_numbers_tables_and_render_gap(app, tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches
    monkeypatch.setattr(knowledge, "office_executable", lambda app: None)
    knowledge.register(app)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "转写数据"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "指标"
    table.cell(1, 0).text = "延迟"
    table.cell(1, 1).text = "3 秒"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "第二页"
    source = tmp_path / "slides.pptx"
    presentation.save(source)
    collection = app.call("knowledge.create", name="幻灯片")
    result = app.call("knowledge.ingest", collection_id=collection["id"], paths=[str(source)], visual=True)
    document = result["result"]["documents"][0]
    assert document["pages"] == 2 and document["warnings"]
    results = app.call("knowledge.search", collection_id=collection["id"], query="延迟")["results"]
    assert any(r["page"] == 1 and "3 秒" in r["text"] for r in results)


def test_knowledge_docx_tables_and_no_invented_pages(app, tmp_path, monkeypatch):
    from docx import Document
    monkeypatch.setattr(knowledge, "office_executable", lambda app: None)
    knowledge.register(app)
    document = Document()
    document.add_paragraph("转写报告")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "吞吐量"
    table.cell(0, 1).text = "42"
    source = tmp_path / "report.docx"
    document.save(source)
    collection = app.call("knowledge.create", name="文档")
    result = app.call("knowledge.ingest", collection_id=collection["id"], paths=[str(source)], visual=False)
    assert result["result"]["warnings"]
    rows = app.call("knowledge.search", collection_id=collection["id"], query="吞吐量")["results"]
    assert rows[0]["page"] is None and "42" in rows[0]["text"]


def make_pdf(path):
    # Minimal two-page PDF fixture; no ReportLab runtime dependency.
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Resources << /Font << /F1 5 0 R >> >> /Contents 6 0 R >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Resources << /Font << /F1 5 0 R >> >> /Contents 7 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for text in (b"Latency three seconds", b"Chart on page two"):
        stream = b"BT /F1 12 Tf 20 250 Td (" + text + b") Tj ET"
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, value in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(str(i).encode() + b" 0 obj\n" + value + b"\nendobj\n")
    position = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{position}\n%%EOF".encode())
    path.write_bytes(output)


def test_knowledge_pdf_every_page_rendered_and_vision_separate(app, tmp_path):
    knowledge.register(app)
    source = tmp_path / "source.pdf"
    make_pdf(source)
    collection = app.call("knowledge.create", name="PDF")
    result = app.call("knowledge.ingest", collection_id=collection["id"], paths=[str(source)], visual=True)
    assert result["result"]["documents"][0]["pages"] == 2
    visions = [messages for role, messages in app.providers.messages if role == "vision"]
    assert len(visions) == 2
    assert visions[0][-1]["content"][1]["image_url"]["url"].startswith("data:image/jpeg;base64,")
    rows = app.call("knowledge.search", collection_id=collection["id"], query="准确率")["results"]
    assert any(r["kind"] == "vision" and Path(r["page_image"]).is_file() for r in rows)


def test_knowledge_image_ingestion_uses_visual_source(app, tmp_path):
    from PIL import Image
    knowledge.register(app)
    source = tmp_path / "chart.png"
    Image.new("RGB", (32, 32), "white").save(source)
    collection = app.call("knowledge.create", name="图像")
    result = app.call("knowledge.ingest", collection_id=collection["id"], paths=[str(source)], visual=True)
    assert result["result"]["documents"][0]["status"] == "ready"
    rows = app.call("knowledge.search", collection_id=collection["id"], query="准确率")["results"]
    assert rows[0]["kind"] == "vision" and rows[0]["page"] == 1


def test_backup_roundtrip_rebases_paths_and_excludes_secrets(app, tmp_path):
    backups.register(app)
    write_json(app.data_dir / "settings.json", {"preferences": {"library_path": str(app.data_dir / "media")}, "roles": {}})
    write_json(app.data_dir / "secrets.json", {"private": "encrypted-test"})
    session = app.data_dir / "sessions" / "one"
    session.mkdir(parents=True)
    (session / "microphone-00001.wav").write_bytes(b"recording")
    write_json(session / "session.json", {"recording": str(session / "microphone-00001.wav")})
    package = tmp_path / "portable.wtbak"
    result = app.call("backups.export", path=str(package), include_media=True)
    assert not result["result"]["encrypted"]
    with zipfile.ZipFile(package) as archive:
        assert "secrets.json" not in archive.namelist()
        assert "sessions/one/microphone-00001.wav" in archive.namelist()
    target = App(tmp_path / "different-directory")
    backups.register(target)
    target.call("backups.import", path=str(package))
    restored = json.loads((target.data_dir / "sessions" / "one" / "session.json").read_text())
    assert restored["recording"].startswith(str(target.data_dir))
    assert target.settings.get()["preferences"]["library_path"] == str(target.data_dir / "media")
    assert (target.data_dir / "sessions" / "one" / "microphone-00001.wav").read_bytes() == b"recording"


def test_backup_password_not_persisted_and_wrong_password_no_mutation(app, tmp_path):
    backups.register(app)
    write_json(app.data_dir / "settings.json", {"preferences": {"test": 1}, "roles": {}})
    package = tmp_path / "encrypted.wtbak"
    app.call("backups.export", path=str(package), include_secrets=True, password="test-password")
    assert "password" not in app.jobs.last_params
    assert "test-password" not in json.dumps(app.jobs.last_params)
    assert package.read_bytes().startswith(backups.MAGIC)
    before = (app.data_dir / "settings.json").read_bytes()
    with pytest.raises(ValueError, match="密码错误"):
        app.call("backups.import", path=str(package), password="incorrect")
    assert (app.data_dir / "settings.json").read_bytes() == before
    app.call("backups.import", path=str(package), password="test-password")
    assert app.settings.secrets["example"] == "secret-value-test-only"


def test_backup_snapshot_sqlite_wal_and_rebase(tmp_path):
    source, target = tmp_path / "source.sqlite3", tmp_path / "copied.sqlite3"
    with contextlib.closing(sqlite3.connect(source)) as db:
        db.execute("PRAGMA journal_mode=WAL")
        db.execute("CREATE TABLE things (path TEXT)")
        db.execute("INSERT INTO things VALUES (?)", (str(tmp_path / "old" / "file.txt"),))
        db.commit()
        backups.snapshot_database(source, target)
    stage = tmp_path / "stage"
    stage.mkdir()
    os.replace(target, stage / target.name)
    backups.rebase_snapshot(stage, tmp_path / "old", tmp_path / "new")
    with contextlib.closing(sqlite3.connect(stage / target.name)) as db:
        assert db.execute("SELECT path FROM things").fetchone()[0] == str(tmp_path / "new" / "file.txt")


def test_backup_rejects_tampered_archive(app, tmp_path):
    backups.register(app)
    write_json(app.data_dir / "settings.json", {"preferences": {"test": "original"}})
    package = tmp_path / "safe.wtbak"
    app.call("backups.export", path=str(package))
    bad = tmp_path / "tampered.wtbak"
    with zipfile.ZipFile(package) as original, zipfile.ZipFile(bad, "w") as archive:
        for name in original.namelist():
            archive.writestr(name, b"modified" if name == "settings.json" else original.read(name))
    before = (app.data_dir / "settings.json").read_bytes()
    with pytest.raises(ValueError, match="校验失败"):
        app.call("backups.import", path=str(bad))
    assert (app.data_dir / "settings.json").read_bytes() == before


def test_backup_restore_rollback_on_reload_failure(app, tmp_path, monkeypatch):
    backups.register(app)
    write_json(app.data_dir / "settings.json", {"preferences": {"test": "archived"}})
    package = tmp_path / "restore.wtbak"
    app.call("backups.export", path=str(package))
    write_json(app.data_dir / "settings.json", {"preferences": {"test": "current"}})
    original_reload = app.settings.reload
    calls = 0
    def fail_once():
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("reload failed")
        original_reload()
    monkeypatch.setattr(app.settings, "reload", fail_once)
    with pytest.raises(RuntimeError, match="reload failed"):
        app.call("backups.import", path=str(package))
    assert app.settings.get()["preferences"]["test"] == "current"
    assert not app.maintenance


def test_backup_external_library_is_portable(app, tmp_path):
    backups.register(app)
    external = tmp_path / "external-disk" / "library"
    external.mkdir(parents=True)
    (external / "talk.mp4").write_bytes(b"media-content")
    write_json(external / "result.json", {"audio": str(external / "talk.mp4")})
    write_json(app.data_dir / "settings.json", {"preferences": {"library_path": str(external)}, "roles": {}})
    app.settings.reload()
    package = tmp_path / "library.wtbak"
    app.call("backups.export", path=str(package), include_media=True)
    target = App(tmp_path / "new-computer")
    backups.register(target)
    target.call("backups.import", path=str(package))
    restored_library = Path(target.settings.get()["preferences"]["library_path"])
    assert restored_library.is_relative_to(target.data_dir)
    assert (restored_library / "talk.mp4").read_bytes() == b"media-content"
    assert json.loads((restored_library / "result.json").read_text())["audio"] == str(restored_library / "talk.mp4")


def test_backup_omitted_recordings_preserved_when_restoring(app, tmp_path):
    backups.register(app)
    session = app.data_dir / "sessions" / "a"
    session.mkdir(parents=True)
    (session / "microphone-00001.wav").write_bytes(b"keep this recording")
    write_json(session / "session.json", {"id": "a"})
    package = tmp_path / "without-media.wtbak"
    app.call("backups.export", path=str(package), include_media=False)
    with zipfile.ZipFile(package) as bundle:
        assert "sessions/a/microphone-00001.wav" not in bundle.namelist()
    app.call("backups.import", path=str(package))
    assert (session / "microphone-00001.wav").read_bytes() == b"keep this recording"


def wait_real_job(app, job_id, timeout=15):
    import time
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        result = app.jobs.get(job_id)
        if result["status"] not in ("running", "queued", "cancelling") and job_id not in app.jobs.active:
            assert result["status"] == "completed", result
            return result
        time.sleep(0.02)
    raise AssertionError("job timed out")


def test_real_app_portable_restore_knowledge_plugins_jobs_and_dpapi(tmp_path, monkeypatch):
    from toolbox.app import App as RealApp
    original = RealApp(tmp_path / "original-app", register_live=False)
    target = RealApp(tmp_path / "restored-app", register_live=False)
    provider = Providers()
    monkeypatch.setattr(original.providers, "chat", provider.chat)
    monkeypatch.setattr(original.providers, "embed", provider.embed)
    monkeypatch.setattr(target.providers, "chat", provider.chat)
    monkeypatch.setattr(target.providers, "embed", provider.embed)
    try:
        original.call("settings.update", {"settings": {"providers": [{"id": "test", "kind": "openai", "name": "Fixture provider", "base_url": "http://127.0.0.1:1", "api_key": "test-fixture-key-not-real"}]}})
        collection = original.call("knowledge.create", {"name": "Test sources"})
        source = tmp_path / "isolated-source.txt"
        source.write_text("转写延迟三秒", encoding="utf-8")
        ingestion = original.call("knowledge.ingest", {"collection_id": collection["id"], "paths": [str(source)], "visual": False})
        wait_real_job(original, ingestion["id"])
        original.call("plugins.install", {"path": str(Path(__file__).parents[1] / "examples/file-hash.toolpkg")})
        outside_artifact = tmp_path / "custom-output" / "summary.md"
        outside_artifact.parent.mkdir()
        outside_artifact.write_text("外部目录中的任务结果", encoding="utf-8")
        def artifact_runner(job):
            job.artifact(outside_artifact, "markdown", "Summary")
            return {"path": str(outside_artifact)}
        external_job = original.jobs.submit("fixture-output", {}, artifact_runner)
        wait_real_job(original, external_job["id"])
        package = tmp_path / "real-portable.wtbak"
        exported = original.call("backups.export", {"path": str(package), "include_secrets": True, "password": "fixture-password"})
        wait_real_job(original, exported["id"])
        imported = target.call("backups.import", {"path": str(package), "password": "fixture-password"})
        wait_real_job(target, imported["id"])
        assert target.settings.secret("test") == "test-fixture-key-not-real"
        assert all("password" not in j["params"] for j in target.jobs.list())
        assert target.call("plugins.list")["plugins"][0]["id"] == "file-hash"
        migrated = target.jobs.get(external_job["id"])["artifacts"][0]["path"]
        assert Path(migrated).is_relative_to(target.data_dir)
        assert Path(migrated).read_text(encoding="utf-8") == "外部目录中的任务结果"
        results = target.call("knowledge.search", {"collection_id": collection["id"], "query": "转写"})["results"]
        assert results and Path(results[0]["source_path"]).is_relative_to(target.data_dir)
        assert Path(results[0]["source_path"]).is_file()
        plugin = target.call("plugins.run", {"id": "file-hash", "params": {"path": results[0]["source_path"]}})
        completed = wait_real_job(target, plugin["id"])
        assert completed["result"]["hash"] == hashlib.sha256(source.read_bytes()).hexdigest()
    finally:
        original.close()
        target.close()
        original.jobs.pool.shutdown(wait=True)
        target.jobs.pool.shutdown(wait=True)
